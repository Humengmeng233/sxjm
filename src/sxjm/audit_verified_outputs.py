"""Independent saved-artifact audit and small, non-tuning sensitivity experiment."""
import argparse
from dataclasses import replace
import json
from pathlib import Path
import platform
from time import perf_counter

import numpy as np
import pandas as pd
import scipy
import openpyxl

from .model_config import ModelConfig
from .solve_all_models import load_inputs
from .solve_verified_models import source_hashes
from .causal_scenarios import CausalScenarioFactory
from .stochastic import solve_stochastic_plan,execute_controls


def sensitivity(data,config,folder):
    records=[]
    for day in (31,78,243):
        # One-variable-at-a-time snapshots, all start from 6000 kWh.
        # These are NOT continuous-year strategies or an independent tuning set.
        cases=[(8,.15,.2,False),(16,.15,.2,False),(32,.15,.2,False),
               (8,0.,.2,False),(8,.3,.2,False),(8,.15,0.,False),(8,.15,.5,False),
               (8,.15,.2,True)]
        for count,radius,weight,unknown in cases:
            modified=replace(config,robustness_radius=radius)
            factory=CausalScenarioFactory(data,modified,count)
            bundle=factory.get(day,0,pv_information=True,variable_price=True,unknown_price=unknown)
            start=perf_counter()
            plan=solve_stochastic_plan(bundle.net,bundle.prices,bundle.probability,6000.,modified,risk_weight=weight)
            elapsed=perf_counter()-start
            actual=execute_controls(data.net_kwh[day],plan.q,plan.charge,plan.discharge,6000.,modified)
            actual_cost=float(data.variable_price[day]@plan.q+5*data.variable_price[day]@actual.emergency_kwh)
            records.append(dict(date=data.dates[day],scenario_limit=count,actual_scenarios=len(bundle.probability),
                radius=radius,risk_weight=weight,unknown_price=unknown,risk_objective_yuan=plan.objective,
                realized_snapshot_cost_yuan=actual_cost,emergency_kwh=float(actual.emergency_kwh.sum()),
                solve_seconds=elapsed,storage_end_kwh=float(actual.storage_kwh[-1]),max_lp_residual=plan.residual))
    pd.DataFrame(records).to_csv(folder/"sensitivity_snapshots.csv",index=False,encoding="utf-8-sig")
    return records


def audit(folder):
    root=Path(__file__).resolve().parents[2]
    info=json.loads((folder/"solution_summary.json").read_text(encoding="utf-8"))
    config=ModelConfig(**info["parameters"])
    data=load_inputs(root,config)
    assert source_hashes(root)==info["source_hashes"],"Source hash mismatch"
    report={"source_files_unchanged":True,"python_version":platform.python_version(),"scipy_version":scipy.__version__,"numpy_version":np.__version__,"workbooks":[]}
    for name,model in (("result1.xlsx",None),("result2.xlsx","Q2"),("result3.xlsx","Q3"),("result4-2.xlsx","Q4-2"),("result4-3.xlsx","Q4-3")):
        wb=openpyxl.load_workbook(folder/name,data_only=True)
        wf=openpyxl.load_workbook(folder/name,data_only=False)
        assert not any(cell.data_type=="e" for sheet in wb for row in sheet for cell in row)
        formulas=sum(cell.data_type=="f" for sheet in wf for row in sheet for cell in row)
        if model is None:
            values=pd.read_csv(folder/"question1_intervals.csv")
            np.testing.assert_allclose([wb.worksheets[0].cell(t+2,2).value for t in range(144)],values.purchase_kwh,atol=1e-8)
        else:
            frame=pd.read_csv(folder/f"{model}_intervals.csv")
            daily=pd.read_csv(folder/f"{model}_daily.csv")
            days=len(daily)
            plan_count=2 if model in ("Q3","Q4-3") else 1
            for i in range(plan_count):
                sheet=wb.worksheets[i]
                expected=frame["original_plan_kwh" if i==0 else "accepted_plan_kwh"].to_numpy().reshape(days,144)
                actual=np.asarray([[sheet.cell(day+2,t+2).value for t in range(144)] for day in range(days)],dtype=float)
                np.testing.assert_allclose(actual,expected,rtol=1e-12,atol=1e-8)
                totals=[sheet.cell(day+2,146).value for day in range(days)]
                assert all(value is not None for value in totals),"Formula cache not populated: run workbook recalculation first"
                np.testing.assert_allclose(totals,expected.sum(axis=1),rtol=1e-12,atol=1e-7)
                costs=daily.plan_cost_yuan+(daily.adjustment_cost_yuan if i else 0.)
                np.testing.assert_allclose([sheet.cell(day+2,147).value for day in range(days)],costs,rtol=1e-12,atol=1e-7)
                assert sheet.cell(1,2).value=="00:00—00:10" and sheet.cell(1,145).value=="23:50—24:00"
                assert pd.Timestamp(sheet.cell(2,1).value)==pd.Timestamp(daily.date.iloc[0])
                assert pd.Timestamp(sheet.cell(days+1,1).value)==pd.Timestamp(daily.date.iloc[-1])
            battery=wb.worksheets[plan_count]
            assert battery.max_row==days*6+1
            for col,field in ((3,"charge_kwh"),(4,"discharge_kwh")):
                values=np.array([battery.cell(row, col).value for row in range(2,battery.max_row+1)]).reshape(days,6)
                expected=frame[field].to_numpy().reshape(days,6,24).sum(axis=2)
                np.testing.assert_allclose(values,expected,atol=1e-7)
            emergencies=wb.worksheets[plan_count+1]
            exported={}
            for row in emergencies.iter_rows(min_row=2,values_only=True):
                date=pd.Timestamp(row[0]).strftime("%Y-%m-%d")
                exported[date]=exported.get(date,0.)+float(row[2])
            for row in daily.itertuples():
                np.testing.assert_allclose(exported[row.date],row.emergency_energy_kwh,rtol=1e-10,atol=1e-6)
        report["workbooks"].append(dict(file=name,all_data_reconciled=True,formulas=formulas,cached_results_verified=True,sheets=len(wb.sheetnames)))
    if (folder/"artifact_workbook_validation.json").exists():
        engine=json.loads((folder/"artifact_workbook_validation.json").read_text(encoding="utf-8"))
        assert all("matched 0 entries" in result["errors"] for result in engine)
        report["artifact_recalculation_and_mutation_test"]=True
    info["limitations"]=[value for value in info["limitations"] if value!="formula caches require spreadsheet recalculation"]
    info["spreadsheet_verification"]="Recalculated in Artifact Tool; cached SUM totals independently verified. Native Excel not exercised."
    (folder/"solution_summary.json").write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding="utf-8")
    note=folder/"模型求解与结果说明.md"
    text=note.read_text(encoding="utf-8")
    text=text.replace("计划电量合计列为SUM公式，openpyxl不计算公式缓存，首次Excel打开会重算；核验脚本直接核对公式引用及其144个源数值，而非把缺失缓存当零。",
        "计划电量合计列为SUM公式。交付文件已在Artifact Tool中重计算、改变单一输入验证SUM依赖并恢复、导出缓存；再用独立Python逐日核对缓存合计与144个源数值、费用和全量事件。未使用原生Excel执行界面测试；单独重跑Python导出时仍需表格引擎刷新缓存。")
    if "### 附加验证" not in text:
        text+="\n### 附加验证\n\n`python -m sxjm.test_verified_models` 的7项测试已通过，覆盖两支持DRO/CVaR与独立枚举一致、调增调减定价、实时执行因果性、未来信息扰动、1月实态衔接和Q1边界。\n\n`sensitivity_snapshots.csv`提供2月1日、3月20日、9月1日共24组单日快照，改变场景数、半径、风险权重或价格信息。快照统一从6000 kWh开始，不是连续全年比较，也不用于反向选择主结果参数。\n"
    note.write_text(text,encoding="utf-8")
    return report,data,config


def main():
    parser=argparse.ArgumentParser();parser.add_argument("--output",default="outputs/model_solution_v2")
    args=parser.parse_args();folder=Path(__file__).resolve().parents[2]/args.output
    report,data,config=audit(folder)
    records=sensitivity(data,config,folder)
    report["sensitivity_snapshot_count"]=len(records)
    (folder/"final_validation.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":main()

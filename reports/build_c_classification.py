"""Export the problem-type classification tables; diagrams remain in Markdown."""

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from build_c_analysis_workbook import add_sheet, markdown_tables


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    folder = Path(__file__).resolve().parent
    report = folder / "C题_题型归类与思维导图.md"
    output = folder / "C题_题型归类.xlsx"
    content = report.read_text(encoding="utf-8")
    tables = markdown_tables(content)
    charts = re.findall(r"```mermaid\n(.*?)\n```", content, re.S)
    assert len(charts) == 2
    for chart in charts:
        assert "flowchart " in chart
        init = re.search(r"%%\{init:\s*(.*?)\}%%", chart)
        assert init is not None
        json.loads(init.group(1))
        nodes = set(re.findall(r"\b(\w+)\s*(?:\[|\(\[)", chart))
        for declaration in re.findall(r"^\s*class\s+([\w,]+)\s+\w+;", chart, re.M):
            assert set(declaration.split(",")) <= nodes, declaration
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "C题题型归类与判定依据"
    add_sheet(wb, "阅读说明", [
        ["项目", "说明"],
        ["来源", str(folder.parent / "C题.docx")],
        ["完整文档与导图", str(report)],
        ["核心判定", "四问均以优化决策类为主；问题3额外明确要求预报必要性评价。"],
        ["层级区别", "主＝最终任务；辅＝辅助分析；基＝机理基础；条件＝依据信息或方法选择；—＝不构成独立题型。"],
        ["边界", "不能因使用预报就要求重训模型，也不能因出现电网就判为网络构建。"],
        ["导图", "Markdown内含完整四问思维导图与递进关系图；颜色保持问题分支一致。"],
        ["结果状态", "题型分析与任务结构说明；未开展购电优化求解。"],
    ], [25, 128])
    expected = [
        ("1.", "六类判定标准", [22, 76, 82, 95], 7),
        ("2.", "分问题型与依据", [17, 41, 90, 94, 99, 101], 5),
        ("3.", "六类题型矩阵", [22, 42, 92, 85, 22, 58, 24], 6),
        ("4.", "16项子任务归类", [37, 72, 123], 17),
        ("6.1", "逻辑接口", [37, 95, 92, 118], 7),
    ]
    for prefix, title, widths, row_count in expected:
        found = [rows for heading, rows in tables if heading.startswith(prefix + " ")]
        assert len(found) == 1
        assert len(found[0]) == row_count, title
        add_sheet(wb, title, found[0], widths)
    wb.save(output)
    wb.close()
    check = load_workbook(output, read_only=True)
    assert check["分问题型与依据"].max_row == 5
    assert check["16项子任务归类"].max_row == 17
    print(json.dumps({"file": str(output), "worksheets": check.sheetnames, "main_questions": 4, "subtasks": 16, "mermaid_diagrams": len(charts), "diagram_checks": "init JSON and declared node references checked; not browser-rendered"}, ensure_ascii=False, indent=2))
    check.close()


if __name__ == "__main__":
    main()

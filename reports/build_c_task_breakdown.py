"""Export the reviewed C-problem task breakdown as an XLSX and one CSV table."""

from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from build_c_analysis_workbook import add_sheet, docx_contents, markdown_tables, verify_source


HERE = Path(__file__).resolve().parent
REPORT = HERE / "C题_分问目标与递进逻辑.md"
OUTPUT = HERE / "C题_分问目标与递进逻辑.xlsx"
CSV_OUTPUT = HERE / "C题_分问目标与递进逻辑__子任务拆分.csv"


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    source_tables = markdown_tables((HERE / "C题_逐句解析.md").read_text(encoding="utf-8"))
    statements = [row for _, rows in source_tables for row in rows[1:] if re.fullmatch(r"S\d{2}", row[0])]
    original_tables = [row for _, rows in source_tables for row in rows[1:] if re.fullmatch(r"T\d{2}", row[0])]
    paragraphs, _ = docx_contents()
    source_verification = verify_source(paragraphs, statements, original_tables)

    content = REPORT.read_text(encoding="utf-8")
    tables = markdown_tables(content)
    by_heading = dict(tables)
    tasks = [row for _, rows in tables for row in rows[1:] if re.fullmatch(r"[1-4]\.\d", row[0])]
    expected = [f"{question}.{part}" for question, count in [(1, 3), (2, 4), (3, 5), (4, 4)] for part in range(1, count + 1)]
    assert [row[0] for row in tasks] == expected, "Missing or duplicated subtask identifiers"
    assert all(len(row) == 6 and all(cell.strip() for cell in row) for row in tasks)
    required_files = ["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"]
    assert all(filename in content for filename in required_files)
    assert len(re.findall(r"(?m)^\$\$$", content)) % 2 == 0
    assert len(re.findall(r"(?m)^```", content)) % 2 == 0

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "C题：分问目标、输出要求与递进逻辑"
    wb.properties.subject = "16项子任务；直接与隐含目标；数值、公式、方案与专业图表；递进关系"
    add_sheet(wb, "阅读说明", [
        ["项目", "说明"],
        ["题目来源", str(HERE.parent / "C题.docx")],
        ["完整说明", str(REPORT)],
        ["拆分范围", "4个大问拆为16项子任务；编号为分析性拆分，不是原题额外小问。"],
        ["目标区别", "直接目标对应题面任务；隐含铺垫目标是结构推断，不是新增硬性要求。"],
        ["要求性质", "明示＝题面规定；支撑＝可信建模与核验所需；建议＝可选择的图形或对照设计。"],
        ["当前状态", "本文件是任务设计和输出清单，尚未求解或生成带结果数值的图表。"],
        ["关键比较", "问题3应设置相同0:00预报下不做日内更新的内部对照；问题4分4-2和4-3两条分支。"],
        ["CSV", str(CSV_OUTPUT)],
        ["依据核对", "上一份逐句解析52条原文单元、4张题表和93个非空段落已再次与DOCX核对。"],
    ], [24, 130])

    header = ["编号", "拆分任务", "直接目标", "隐含铺垫目标（推断）", "输出要求", "要求性质"]
    task_rows = [header] + tasks
    add_sheet(wb, "16项子任务", task_rows, [10, 37, 80, 77, 96, 58])
    wb["16项子任务"].freeze_panes = "C2"

    targets = [
        ("1.", "四问目标总览", [16, 86, 89, 65, 70]),
        ("6.", "输出要求矩阵", [16, 90, 79, 79, 75, 60]),
        ("6.2", "文件工作表", [27, 57, 91, 80]),
        ("7.", "专业图表建议", [10, 30, 47, 100, 90]),
        ("8.1", "递进依赖关系", [49, 83, 81, 85]),
        ("8.2", "电价调度二维关系", [47, 57, 67]),
        ("9.1", "比较设计", [37, 102, 112, 104]),
        ("10.", "完成判据", [16, 110, 100]),
    ]
    for prefix, title, widths in targets:
        found = [rows for heading, rows in by_heading.items() if heading.startswith(prefix + " ")]
        if len(found) != 1:
            raise ValueError(f"Expected one table for {title}; found {len(found)}")
        add_sheet(wb, title, found[0], widths)
    wb.save(OUTPUT)
    wb.close()

    with CSV_OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows(task_rows)

    check = load_workbook(OUTPUT, read_only=True, data_only=False)
    actual = list(check["16项子任务"].values)
    assert actual == [tuple(row) for row in task_rows]
    assert check["输出要求矩阵"].max_row == 6
    assert check["文件工作表"].max_row == 6
    assert check["专业图表建议"].max_row == 9
    assert check["递进依赖关系"].max_row == 8
    names = check.sheetnames
    check.close()
    with CSV_OUTPUT.open(encoding="utf-8-sig", newline="") as handle:
        assert list(csv.reader(handle)) == task_rows
    print(json.dumps({
        "source_statements_verified": source_verification["source_statements_verified"],
        "subtasks": len(tasks),
        "subtasks_by_question": {str(question): sum(row[0].startswith(f"{question}.") for row in tasks) for question in range(1, 5)},
        "all_five_result_files_covered": True,
        "workbook_sheets": names,
        "output_workbook": str(OUTPUT),
        "output_csv": str(CSV_OUTPUT),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

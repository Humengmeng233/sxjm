from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CSV_DIR = ROOT / "data" / "csv" / "raw-datta"
XLSX_DIR = ROOT / "data" / "raw-datta"


def numeric_summary(frame: pd.DataFrame, columns: list[str]) -> dict[str, object]:
    values = frame[columns].apply(pd.to_numeric, errors="coerce")
    arr = values.to_numpy(dtype=float)
    finite = np.isfinite(arr)
    return {
        "cells": int(arr.size),
        "missing": int(np.isnan(arr).sum()),
        "nonfinite": int((~finite & ~np.isnan(arr)).sum()),
        "negative": int((arr < 0).sum()),
        "zero": int((arr == 0).sum()),
        "min": float(np.nanmin(arr)),
        "max": float(np.nanmax(arr)),
        "mean": float(np.nanmean(arr)),
    }


def main() -> None:
    report: dict[str, object] = {"xlsx": {}, "csv": {}, "cross_checks": {}}

    for path in sorted(XLSX_DIR.glob("*.xlsx")):
        wb = load_workbook(path, read_only=False, data_only=False)
        report["xlsx"][path.name] = {
            "sheets": [
                {
                    "name": ws.title,
                    "rows": ws.max_row,
                    "columns": ws.max_column,
                    "merged_ranges": len(ws.merged_cells.ranges),
                    "formulas": sum(
                        1
                        for row in ws.iter_rows()
                        for cell in row
                        if isinstance(cell.value, str) and cell.value.startswith("=")
                    ),
                }
                for ws in wb.worksheets
            ]
        }

    a1 = pd.read_csv(CSV_DIR / "附件1__Sheet1.csv")
    a2_pv = pd.read_csv(CSV_DIR / "附件2__光伏发电实际功率.csv")
    a2_load = pd.read_csv(CSV_DIR / "附件2__小区负载.csv")
    a3 = pd.read_csv(CSV_DIR / "附件3__Sheet1.csv")
    a4 = pd.read_csv(CSV_DIR / "附件4__Sheet1.csv")

    report["csv"]["附件1"] = {
        "shape": list(a1.shape),
        "columns": list(a1.columns),
        "dtypes": {k: str(v) for k, v in a1.dtypes.items()},
        "missing_by_column": {k: int(v) for k, v in a1.isna().sum().items()},
        "duplicate_rows": int(a1.duplicated().sum()),
        "duplicate_times": int(a1["时间"].duplicated().sum()),
        "time_first_last": [str(a1["时间"].iloc[0]), str(a1["时间"].iloc[-1])],
        "numeric": {
            c: numeric_summary(a1, [c]) for c in ["电价", "小区负载", "光伏发电预测功率"]
        },
    }

    for name, frame in [("附件2_负载", a2_load), ("附件2_光伏实际", a2_pv), ("附件4_电价", a4)]:
        date_col = frame.columns[0]
        dates = pd.to_datetime(frame[date_col], errors="coerce")
        numeric_cols = list(frame.columns[1:])
        report["csv"][name] = {
            "shape": list(frame.shape),
            "date_column": date_col,
            "date_dtype_raw": str(frame[date_col].dtype),
            "value_dtypes": sorted({str(v) for v in frame[numeric_cols].dtypes}),
            "date_first_last": [str(dates.min().date()), str(dates.max().date())],
            "date_parse_failures": int(dates.isna().sum()),
            "duplicate_dates": int(dates.duplicated().sum()),
            "missing_calendar_days": int(
                len(pd.date_range(dates.min(), dates.max(), freq="D").difference(pd.DatetimeIndex(dates)))
            ),
            "time_columns": len(numeric_cols),
            "time_first_last": [numeric_cols[0], numeric_cols[-1]],
            "duplicate_time_headers": int(pd.Index(numeric_cols).duplicated().sum()),
            "numeric": numeric_summary(frame, numeric_cols),
            "rows_with_any_missing": int(frame[numeric_cols].isna().any(axis=1).sum()),
        }

    forecast_cols = [c for c in a3.columns if c.startswith("预报") and c.endswith("小时")]
    filled_dates = pd.to_datetime(a3["日期"].ffill(), errors="coerce")
    release = a3["预报时刻"].astype(str)
    report["csv"]["附件3"] = {
        "shape": list(a3.shape),
        "columns": list(a3.columns),
        "dtypes": {k: str(v) for k, v in a3.dtypes.items()},
        "raw_blank_dates": int(a3["日期"].isna().sum()),
        "filled_date_first_last": [str(filled_dates.min().date()), str(filled_dates.max().date())],
        "date_parse_failures_after_fill": int(filled_dates.isna().sum()),
        "unique_filled_dates": int(filled_dates.nunique()),
        "duplicate_date_release_pairs_after_fill": int(
            pd.DataFrame({"date": filled_dates, "release": release}).duplicated().sum()
        ),
        "release_counts": {str(k): int(v) for k, v in release.value_counts().sort_index().items()},
        "forecast_columns": len(forecast_cols),
        "numeric": numeric_summary(a3, forecast_cols),
        "rows_with_any_missing_forecast": int(a3[forecast_cols].isna().any(axis=1).sum()),
    }

    report["cross_checks"] = {
        "attachment2_shapes_match": bool(a2_load.shape == a2_pv.shape),
        "attachment2_headers_match": bool(list(a2_load.columns) == list(a2_pv.columns)),
        "attachment2_dates_match": bool(a2_load.iloc[:, 0].equals(a2_pv.iloc[:, 0])),
        "attachment4_shape_matches_attachment2": bool(a4.shape == a2_load.shape),
        "attachment4_headers_match_attachment2": bool(list(a4.columns) == list(a2_load.columns)),
        "attachment4_dates_match_attachment2": bool(a4.iloc[:, 0].equals(a2_load.iloc[:, 0])),
        "forecast_has_four_releases_per_day": bool(
            pd.DataFrame({"date": filled_dates, "release": release})
            .groupby("date")["release"]
            .nunique()
            .eq(4)
            .all()
        ),
        "forecast_release_set": sorted(release.unique().tolist()),
    }

    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
